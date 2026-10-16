#!/usr/bin/env python3
"""
Robô de leitura de Notas Fiscais escaneadas
===========================================

Lê os arquivos escaneados de uma pasta (PDF, JPG, PNG, TIF...), faz OCR em
cada página e monta uma planilha Excel com:

    Fornecedor | Nº NF | Pedido (45...) | + colunas para conferência

Uso:
    python robo_nf.py                          # lê a pasta "notas" ao lado do script
    python robo_nf.py "C:\\Scans\\Notas"        # lê outra pasta
    python robo_nf.py notas -o outubro.xlsx    # escolhe o nome da planilha
    python robo_nf.py notas --salvar-texto     # salva o texto lido de cada página
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import shutil
import sys
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from pathlib import Path

# ----------------------------------------------------------------------------
# Configurações
# ----------------------------------------------------------------------------

# Pedido de compra: começa com "45" e tem 10 dígitos (padrão SAP, ex.: 4500012345).
# Se os seus pedidos tiverem outro tamanho, é só ajustar aqui.
PREFIXO_PEDIDO = "45"
DIGITOS_PEDIDO = 10

PASTA_SCRIPT = Path(__file__).resolve().parent
PASTA_PADRAO = PASTA_SCRIPT / "notas"
SAIDA_PADRAO = PASTA_SCRIPT / "planilha_nfs.xlsx"
CADASTRO_PADRAO = PASTA_SCRIPT / "fornecedores.csv"  # opcional: CNPJ;Fornecedor

EXTENSOES = {".pdf", ".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}
DPI_PDF = 300

CODIGOS_UF = {11, 12, 13, 14, 15, 16, 17, 21, 22, 23, 24, 25, 26, 27, 28, 29,
              31, 32, 33, 35, 41, 42, 43, 50, 51, 52, 53}

# ----------------------------------------------------------------------------
# Expressões usadas na leitura do texto
# ----------------------------------------------------------------------------

RE_PEDIDO = re.compile(
    rf"(?<!\d){PREFIXO_PEDIDO}\d{{{DIGITOS_PEDIDO - len(PREFIXO_PEDIDO)}}}(?!\d)")
# "Pedido: 4500 012345", "PED. 45000I2345", "Ordem de compra nº ..." (números quebrados pelo OCR)
RE_PEDIDO_ROTULO = re.compile(
    r"(?:PEDIDO|\bPED\b|\bP\.?\s?O\b|XPED|ORDEM\s+DE\s+COMPRA|\bO\.?C\b)"
    r"[^\d\n]{0,25}?([0-9OQDIL|][0-9OQDIL| .\-/]{8,20})")

# Chave de acesso impressa em 11 grupos de 4 (aceita CNPJ alfanumérico)
RE_CHAVE_AGRUPADA = re.compile(r"(?<![0-9A-Z])(?:[0-9A-Z]{4}[ .]{0,2}){10}[0-9A-Z]{4}(?![0-9A-Z])")
RE_SEQ_DIGITOS = re.compile(r"\d(?:[ .]{0,2}\d){43,}")

RE_ROTULO_CHAVE = re.compile(r"CHAVE\s+DE\s+ACESSO")

# Número da NF no DANFE: "Nº 000.012.345" (o OCR às vezes põe espaço junto dos pontos)
_NF_PONTUADO = r"\d{3}\s?[.,]\s?\d{3}\s?[.,]\s?\d{3}"
_NF_ESPACADO = r"\d{3}\s?[.,\s]\s?\d{3}\s?[.,\s]\s?\d{3}"
RE_NF_DANFE_ROTULO = re.compile(
    rf"\bN\s*[O0°Q.:\-]{{0,3}}\s*[:.]?\s*({_NF_ESPACADO})(?!\d|[.,/-]\d)")
# "000093269" logo antes de "SÉRIE" (sem o "Nº")
RE_NF_ZEROS_SERIE = re.compile(r"(?<![\d.,])(0\d{5,8})\s*[-|]?\s*\n?\s*SERIE\b")
# "Nº 93269  SÉRIE 1" / "Nº 000093269" e na linha de baixo "SÉRIE 001"
RE_NF_ANTES_SERIE = re.compile(
    r"\bN\s*[O0°.:\-]{0,3}\s*[:.]?\s*(\d{1,9}|\d{1,3}(?:\.\d{3}){1,2})\s*[-|]?\s*\n?\s*SERIE\b")
RE_NF_DANFE = re.compile(rf"(?<![\d./,-])({_NF_PONTUADO})(?!\d|[.,/-]\d)")
# A inscrição estadual de SC (e de outros estados) tem o mesmo formato 999.999.999
RE_CONTEXTO_IE = re.compile(r"INSCRI|\bI\.?\s?E\b")
RE_NF_OUTROS = [  # outros layouts (NFS-e, notas antigas...)
    re.compile(r"NUMERO\s+DA\s+(?:NOTA|NFS?-?E|NF)\b[^\d\n]{0,25}(\d[\d.]{0,10})"),
    re.compile(r"\bN[O0°.]*\s*(?:DA\s+)?(?:NOTA|NFS?-?E|NF)\b[^\d\n]{0,10}(\d[\d.]{0,10})"),
    re.compile(r"NOTA\s+FISCAL[^\d\n]{0,30}?\bN[O0°.:]+\s*(\d[\d.]{0,10})"),
]
# Valor total da NF: "VALOR TOTAL DA NOTA" (o valor fica logo abaixo, no DANFE)
_DINHEIRO = r"(?:\d{1,3}(?:\.\d{3})+|\d+),\d{2}"
RE_DINHEIRO = re.compile(rf"(?<![\d.,])({_DINHEIRO})(?![\d,%])")
RE_DINHEIRO_PALAVRA = re.compile(rf"(?:R\$)?({_DINHEIRO})")
RE_ROTULO_TOTAL = re.compile(r"(?:\b(?:VALOR|VLR|V)\b\.?\s*)?\bTOTAL\.?\s*(?:DA\s*)?(?:NOTA|NF-?E?)\b")
RE_TOTAL_MESMA_LINHA = re.compile(
    RE_ROTULO_TOTAL.pattern + rf"[ \t:=|]*(?:R\s?\$)?[ \t]*({_DINHEIRO})(?![\d,%])")
RE_LETRAS = re.compile(r"[A-Z]{2}")

RE_SERIE = re.compile(r"\bSERIE\s*[:.]?\s*(\d{1,3})(?!\d)")

RE_CNPJ = re.compile(
    r"(?<![0-9A-Z])([0-9A-Z]{2})\.?([0-9A-Z]{3})\.?([0-9A-Z]{3})\s?/\s?([0-9A-Z]{4})\s?-?\s?(\d{2})(?!\d)")

RE_PALAVRA = re.compile(r"\S+")
FIM_DO_CANHOTO = ("PRODUTOS", "PRODUTO", "SERVICOS", "MERCADORIAS", "MERCADORIA", "ITENS")
RE_EMITENTE = re.compile(r"IDENTIFICA\w*\s+DO\s+EMITENTE")
RE_PRESTADOR = re.compile(r"PRESTADOR")
RE_RAZAO_SOCIAL = re.compile(r"(?:NOME\s*/\s*)?RAZAO\s+SOCIAL\s*[:.]?\s*")
RE_FIM_EMITENTE = re.compile(r"DESTINAT|TOMADOR")
RE_SUFIXO = re.compile(r"\b(?:LTDA|EIRELI|EPP|ME|S[./]?\s?A)\b\.?")
RE_SUFIXO_FINAL = re.compile(r"\b(?:LTDA|EIRELI|EPP|ME|S[./]?\s?A)\b\.?\s*$")
RE_TOKEN_LEGIVEL = re.compile(r"[\w&.,'/()\-]+")
CONECTORES = {"DE", "DO", "DA", "DOS", "DAS", "E", "EM"}
SUFIXOS_EMPRESA = {"LTDA", "LTD", "LT", "SA", "ME", "EPP", "EIRELI", "CIA"}
RE_ENDERECO = re.compile(r"^\s*(?:RUA|AV|AVENIDA|RODOVIA|ROD|ESTRADA|ALAMEDA|TRAVESSA|PRACA|LARGO)\b")
RE_CIDADE_UF = re.compile(r"[-/]\s*(?:AC|AL|AP|AM|BA|CE|DF|ES|GO|MA|MT|MS|MG|PA|PB|PR|PE|PI|RJ|RN|RS|RO|RR|SC|SP|SE|TO)\s*$")
RE_ROTULOS = re.compile(
    r"\b(?:DANFE|DOCUMENTO\s+AUXILIAR|NOTA\s+FISCAL|CNPJ|INSCRICAO|CHAVE\s+DE\s+ACESSO|"
    r"ENTRADA|SAIDA|NATUREZA|IDENTIFICACAO|PROTOCOLO|FOLHA|SERIE|ENDERECO|FONE|CEP|"
    r"DESTINAT\w*|REMETENTE|CODIGO|DESCRICAO|NCM|CFOP|QTDE?|QUANTIDADE|VALOR\s+TOTAL|"
    r"DADOS\s+(?:DOS?|ADICIONAIS)|CALCULO\s+DO\s+IMPOSTO|TRANSPORTADOR\w*|FATURA|DUPLICATAS?)\b")

# Letras que o OCR costuma confundir com números
_TROCA_DIGITOS = str.maketrans({"O": "0", "Q": "0", "D": "0", "I": "1", "L": "1",
                                "|": "1", "S": "5", "B": "8", "Z": "2"})
_LETRAS_CONFUSAS = set("OQDILSBZ|")
RE_TOKEN_NUMERICO = re.compile(r"[0-9A-Z|][0-9A-Z|.,/\-]*")
RE_LETRA_ENTRE_DIGITOS = re.compile(
    r"(?<=\d)[OQDILS|B]{1,3}(?=[\d.])|(?<=[\d.])[OQDILS|B]{1,3}(?=\d)")


# ----------------------------------------------------------------------------
# Dados extraídos
# ----------------------------------------------------------------------------

@dataclass
class DadosNF:
    fornecedor: str = ""
    cnpj: str = ""
    numero_nf: str = ""
    serie: str = ""
    pedidos: list[str] = field(default_factory=list)
    chave: str = ""
    origem_nf: str = ""  # "chave", "danfe", "texto", "chave_parcial" ou "divergente"
    valor_total: str = ""  # "1234.56"
    origem_valor: str = ""  # "rotulo" (junto de "VALOR TOTAL DA NOTA") ou "maior" (estimado)

    @property
    def completo(self) -> bool:
        return bool(self.fornecedor and self.numero_nf and self.pedidos and self.origem_valor == "rotulo")

    @property
    def vazio(self) -> bool:
        return not (self.fornecedor or self.numero_nf or self.pedidos or self.chave)

    def completar_com(self, outro: DadosNF) -> None:
        """Preenche o que estiver faltando com os dados de outra página da mesma NF."""
        for campo in ("fornecedor", "cnpj", "numero_nf", "serie", "chave", "origem_nf",
                      "valor_total", "origem_valor"):
            if not getattr(self, campo):
                setattr(self, campo, getattr(outro, campo))
        if outro.origem_valor == "rotulo" and self.origem_valor != "rotulo":
            self.valor_total, self.origem_valor = outro.valor_total, outro.origem_valor
        for pedido in outro.pedidos:
            if pedido not in self.pedidos:
                self.pedidos.append(pedido)


@dataclass
class RegistroNF:
    arquivo: str
    paginas: list[int]
    dados: DadosNF
    avisos: list[str] = field(default_factory=list)


# ----------------------------------------------------------------------------
# Leitura do texto (funções puras, testáveis sem OCR)
# ----------------------------------------------------------------------------

def normalizar(texto: str) -> str:
    """Maiúsculas e sem acentos, mantendo cada caractere na mesma posição."""
    saida = []
    for c in texto:
        base = "".join(x for x in unicodedata.normalize("NFKD", c) if not unicodedata.combining(x))
        base = base.upper()
        saida.append(base if len(base) == 1 else c)
    return "".join(saida)


def _corrigir_token(m: re.Match) -> str:
    """Token quase todo numérico com letras parecidas com números: OOO.O93.269 -> 000.093.269."""
    token = m.group()
    digitos = sum(c.isdigit() for c in token)
    letras = [c for c in token if c.isalpha() or c == "|"]
    if digitos >= 3 and letras and digitos >= len(letras) and all(c in _LETRAS_CONFUSAS for c in letras):
        return token.translate(_TROCA_DIGITOS)
    return token


def corrigir_digitos(texto: str) -> str:
    """Troca letras confundidas pelo OCR em números (OOO.O93.269, 45000I2345)."""
    texto = RE_TOKEN_NUMERICO.sub(_corrigir_token, texto)
    for _ in range(3):
        novo = RE_LETRA_ENTRE_DIGITOS.sub(lambda m: m.group().translate(_TROCA_DIGITOS), texto)
        if novo == texto:
            break
        texto = novo
    return texto


def _valor(c: str) -> int:
    return ord(c) - 48  # dígitos 0-9 e letras do CNPJ alfanumérico


def dv_mod11(base: str, pesos: list[int]) -> int:
    resto = sum(_valor(c) * p for c, p in zip(base, pesos)) % 11
    return 0 if resto < 2 else 11 - resto


def chave_plausivel(chave: str) -> bool:
    """Confere estrutura e dígito verificador da chave de acesso (44 posições)."""
    if not re.fullmatch(r"\d{6}[0-9A-Z]{12}\d{26}", chave or ""):
        return False
    if int(chave[:2]) not in CODIGOS_UF or not 1 <= int(chave[4:6]) <= 12:
        return False
    if chave[20:22] not in ("55", "65"):  # 55 = NF-e, 65 = NFC-e
        return False
    pesos = [2, 3, 4, 5, 6, 7, 8, 9] * 6
    return dv_mod11(chave[42::-1], pesos) == int(chave[43])


def chave_estrutura_ok(chave: str) -> bool:
    """Chave com estrutura e CNPJ válidos, sem conferir o último dígito (o OCR errou algum número)."""
    return (bool(re.fullmatch(r"\d{6}[0-9A-Z]{12}\d{26}", chave or ""))
            and int(chave[:2]) in CODIGOS_UF and 1 <= int(chave[4:6]) <= 12
            and chave[20:22] in ("55", "65") and cnpj_valido(chave[6:20]))


def cnpj_valido(cnpj: str) -> bool:
    if not re.fullmatch(r"[0-9A-Z]{12}\d{2}", cnpj) or len(set(cnpj)) == 1:
        return False
    pesos = [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    return (dv_mod11(cnpj[:12], pesos[1:]) == int(cnpj[12])
            and dv_mod11(cnpj[:13], pesos) == int(cnpj[13]))


def formatar_cnpj(cnpj: str) -> str:
    if len(cnpj) != 14:
        return cnpj
    return f"{cnpj[:2]}.{cnpj[2:5]}.{cnpj[5:8]}/{cnpj[8:12]}-{cnpj[12:]}"


def encontrar_chave(texto: str, validar=chave_plausivel) -> str:
    for m in RE_CHAVE_AGRUPADA.finditer(texto):
        chave = re.sub(r"[ .]", "", m.group())
        if validar(chave):
            return chave
    trechos = [m.group() for m in RE_SEQ_DIGITOS.finditer(texto)]
    # chave quebrada em mais de uma linha: junta os dígitos logo depois de "CHAVE DE ACESSO"
    trechos += [texto[m.end():m.end() + 250] for m in RE_ROTULO_CHAVE.finditer(texto)]
    for trecho in trechos:
        digitos = re.sub(r"\D", "", trecho)
        for i in range(len(digitos) - 43):
            if validar(digitos[i:i + 44]):
                return digitos[i:i + 44]
    return ""


def _so_numero(trecho: str) -> str:
    digitos = re.sub(r"\D", "", trecho)
    return str(int(digitos)) if digitos and len(digitos) <= 9 and int(digitos) else ""


def _perto_de_ie(texto: str, inicio: int) -> bool:
    return bool(RE_CONTEXTO_IE.search(texto[max(0, inicio - 40):inicio]))


def encontrar_numero_nf(texto: str) -> tuple[str, str]:
    """Devolve (número, origem). Texto já normalizado e com dígitos corrigidos."""
    for regex in (RE_NF_DANFE_ROTULO, RE_NF_ANTES_SERIE, RE_NF_ZEROS_SERIE):
        for m in regex.finditer(texto):
            numero = _so_numero(m.group(1))
            if numero and not _perto_de_ie(texto, m.start()):
                return numero, "danfe"
    # Sem o "Nº" na frente: só números com zeros à esquerda (000.012.345), que não são a IE
    candidatos = Counter(_so_numero(m.group(1)) for m in RE_NF_DANFE.finditer(texto)
                         if m.group(1).startswith("0") and not _perto_de_ie(texto, m.start()))
    candidatos.pop("", None)
    if candidatos:
        return candidatos.most_common(1)[0][0], "danfe"
    for regex in RE_NF_OUTROS:
        for m in regex.finditer(texto):
            numero = _so_numero(m.group(1))
            if numero:
                return numero, "texto"
    return "", ""


def _valor_decimal(valor: str) -> str:
    return valor.replace(".", "").replace(",", ".")  # "1.234,56" -> "1234.56"


def encontrar_valor_total(texto: str) -> tuple[str, str]:
    """Valor total da NF pelo texto. Devolve (valor, origem)."""
    for m in RE_TOTAL_MESMA_LINHA.finditer(texto):
        valor = _valor_decimal(m.group(1))
        if Decimal(valor) > 0:
            return valor, "rotulo"
    # Rótulo no fim da linha e os valores na linha de baixo: o último valor é o da coluna dele
    linhas = texto.split("\n")
    for i, linha in enumerate(linhas):
        for m in RE_ROTULO_TOTAL.finditer(linha):
            if RE_LETRAS.search(linha[m.end():]):
                continue  # há outros rótulos à direita: não dá para saber a coluna
            for prox in linhas[i + 1:i + 3]:
                valores = [_valor_decimal(v) for v in RE_DINHEIRO.findall(prox)]
                if valores:
                    if Decimal(valores[-1]) > 0:
                        return valores[-1], "rotulo"
                    break
    # Sem o rótulo: o maior valor da nota (precisa de conferência)
    valores = [v for v in (_valor_decimal(x) for x in RE_DINHEIRO.findall(texto)) if Decimal(v) > 0]
    if valores:
        return max(valores, key=Decimal), "maior"
    return "", ""


def _inclinacao(palavras: list) -> float:
    """Inclinação da folha (scanner torto), medida entre palavras vizinhas da mesma linha do OCR."""
    inclinacoes = []
    for a, b in zip(palavras, palavras[1:]):
        if len(a) > 5 and len(b) > 5 and a[5] == b[5]:
            dx = (b[1] + b[3] / 2) - (a[1] + a[3] / 2)
            if dx > 0:
                inclinacao = ((b[2] + b[4] / 2) - (a[2] + a[4] / 2)) / dx
                if abs(inclinacao) <= 0.1:
                    inclinacoes.append(inclinacao)
    if not inclinacoes:
        return 0.0
    inclinacoes.sort()
    return inclinacoes[(len(inclinacoes) - 1) // 2]


def linhas_visuais(palavras: list) -> list[list[int]]:
    """Agrupa as palavras [texto, x, y, largura, altura, linha do OCR] nas linhas em que aparecem
    na folha (descontando a inclinação de uma folha escaneada torta)."""
    inclinacao = _inclinacao(palavras)

    def centro(i: int) -> float:
        p = palavras[i]
        return p[2] + p[4] / 2 - inclinacao * (p[1] + p[3] / 2)

    ordem = sorted(range(len(palavras)), key=lambda i: (centro(i), palavras[i][1], i))
    linhas: list = []  # [soma dos centros, altura, índices]: compara com o centro médio da linha
    for i in ordem:
        altura = palavras[i][4]
        if linhas and abs(centro(i) - linhas[-1][0] / len(linhas[-1][2])) <= 0.5 * max(altura, linhas[-1][1]):
            linhas[-1][0] += centro(i)
            linhas[-1][2].append(i)
        else:
            linhas.append([centro(i), altura, [i]])
    return [sorted(linha[2], key=lambda i: (palavras[i][1], i)) for linha in linhas]


def texto_visual(palavras: list) -> str:
    """Texto montado pela posição das palavras: cada rótulo fica junto do seu número/valor."""
    saida = []
    for linha in linhas_visuais(palavras):
        partes, anterior = [], None
        for i in linha:
            x, altura = palavras[i][1], palavras[i][4]
            if anterior is not None:
                xa, la, aa = palavras[anterior][1], palavras[anterior][3], palavras[anterior][4]
                partes.append(" | " if x - (xa + la) > 2 * max(altura, aa) else " ")
            partes.append(palavras[i][0])
            anterior = i
        saida.append("".join(partes))
    return "\n".join(saida)


def valor_pelo_rotulo(palavras: list) -> str:
    """Valor logo abaixo (ou à direita) de "VALOR TOTAL DA NOTA", pela posição na folha."""
    normalizadas = [corrigir_digitos(normalizar(p[0])) for p in palavras]
    dinheiro = [i for i, t in enumerate(normalizadas) if RE_DINHEIRO_PALAVRA.fullmatch(t)]
    melhor = None  # (pontos, índice)
    for linha in linhas_visuais(palavras):
        texto, inicios = "", []
        for i in linha:
            if texto:
                texto += " "
            inicios.append(len(texto))
            texto += normalizadas[i]
        for m in RE_ROTULO_TOTAL.finditer(texto):
            rotulo = [i for i, ini in zip(linha, inicios)
                      if ini < m.end() and ini + len(normalizadas[i]) > m.start()]
            x1 = min(palavras[i][1] for i in rotulo)
            y1 = min(palavras[i][2] for i in rotulo)
            x2 = max(palavras[i][1] + palavras[i][3] for i in rotulo)
            y2 = max(palavras[i][2] + palavras[i][4] for i in rotulo)
            altura = max(palavras[i][4] for i in rotulo)
            cx, cy = (x1 + x2) / 2, (y1 + y2) / 2
            for j in dinheiro:
                x, y, largura, alt = palavras[j][1:5]
                jx, jy = x + largura / 2, y + alt / 2
                # o valor fica dentro da caixa do rótulo: nunca começa antes dele
                abaixo = (y >= y1 + 0.5 * altura and jy - y2 <= 4 * altura
                          and x >= x1 - altura and x <= x2 + 6 * altura)
                ao_lado = abs(jy - cy) <= 0.5 * altura and x >= x2 and x - x2 <= 15 * altura
                if abaixo or ao_lado:
                    pontos = abs(jy - cy) + 0.3 * abs(jx - cx)
                    if melhor is None or pontos < melhor[0]:
                        melhor = (pontos, j)
    if melhor is None:
        return ""
    valor = _valor_decimal(RE_DINHEIRO_PALAVRA.fullmatch(normalizadas[melhor[1]]).group(1))
    return valor if Decimal(valor) > 0 else ""


def encontrar_serie(texto: str) -> str:
    m = RE_SERIE.search(texto)
    return str(int(m.group(1))) if m else ""


def encontrar_cnpj_emitente(norm: str) -> str:
    """Primeiro CNPJ válido antes do bloco do destinatário/tomador."""
    fim = RE_FIM_EMITENTE.search(norm)
    trecho = norm[:fim.start()] if fim else norm
    for m in RE_CNPJ.finditer(trecho):
        bruto = "".join(m.groups())
        for cnpj in (bruto, bruto.translate(_TROCA_DIGITOS)):
            if cnpj_valido(cnpj):
                return cnpj
    return ""


def _juntar_pedido(trecho: str) -> str:
    """'4500 0I2345' -> '4500012345' (só aceita se fechar exatamente um pedido)."""
    partes = re.split(r"[ .\-/]+", trecho.translate(_TROCA_DIGITOS).strip())
    for inicio, parte in enumerate(partes):
        if not parte.startswith(PREFIXO_PEDIDO):
            continue
        digitos = ""
        for pedaco in partes[inicio:]:
            if not pedaco.isdigit():
                break
            digitos += pedaco
            if len(digitos) >= DIGITOS_PEDIDO:
                break
        if RE_PEDIDO.fullmatch(digitos):
            return digitos
    return ""


def encontrar_pedidos(texto: str, ignorar: str = "") -> list[str]:
    achados = [(m.start(), m.group()) for m in RE_PEDIDO.finditer(texto)]
    for m in RE_PEDIDO_ROTULO.finditer(texto):
        pedido = _juntar_pedido(m.group(1))
        if pedido:
            achados.append((m.start(1), pedido))
    pedidos: list[str] = []
    for _, pedido in sorted(achados):
        if pedido not in pedidos and pedido not in ignorar:  # ignora trechos da chave
            pedidos.append(pedido)
    return pedidos


def _limpar_nome(nome: str) -> str:
    nome = re.sub(r"\s+", " ", nome).strip(" ,:;-_|/\\'\"*[](){}").lstrip(".")
    nome = re.sub(r"(?<!\w)\.+$", "", nome).strip()  # mantém o ponto de "S.A." e "LTDA."
    return nome[:100]


def distancia(a: str, b: str) -> int:
    """Quantas letras é preciso trocar, tirar ou pôr para ir de a até b (Levenshtein)."""
    anterior = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        atual = [i]
        for j, cb in enumerate(b, 1):
            atual.append(min(anterior[j] + 1, atual[j - 1] + 1, anterior[j - 1] + (ca != cb)))
        anterior = atual
    return anterior[-1]


def _parece(palavra: str, alvo: str, maximo: int = 2) -> bool:
    return abs(len(palavra) - len(alvo)) <= maximo and distancia(palavra, alvo) <= maximo


def _parece_recebemos(token: str) -> bool:
    return _parece(re.sub(r"[^A-Z]", "", normalizar(token)), "RECEBEMOS")


def _palavra_boa(token: str) -> bool:
    """Palavra com cara de nome: MAIÚSCULA, Capitalizada ou conector (de, da, e...)."""
    for parte in re.split(r"[.,'/()\-&]+", token):
        if not parte or parte.isdigit():
            continue
        if not parte.isalpha():  # letras e números misturados: "st501", "UN00Nd"
            return False
        if not (parte.isupper() or (parte[0].isupper() and parte[1:].islower())
                or (parte.islower() and parte.upper() in CONECTORES)):
            return False
    return True


def _nome_plausivel(nome: str) -> str:
    """Tira pedaços ilegíveis e descarta o nome se ele for "lixo" do OCR."""
    todos = nome.split()
    while todos and (todos[0].upper() == "DE" or _parece_recebemos(todos[0])):
        todos = todos[1:]  # sobra do canhoto: "RECEBEMOS DE ..."
    while todos and (todos[-1].upper() in CONECTORES or not _palavra_boa(todos[-1])):
        todos = todos[:-1]  # nome não termina em "da" nem em lixo
    nome = " ".join(todos)
    tokens = [t for t in todos if RE_TOKEN_LEGIVEL.fullmatch(t)]
    boas = [t for t in tokens if _palavra_boa(t)]
    principais = [t for t in boas if sum(c.isalpha() for c in t) >= 3
                  and re.sub(r"\W", "", t).upper() not in SUFIXOS_EMPRESA | CONECTORES]
    com_numero = sum(1 for t in tokens if any(c.isdigit() for c in t))
    norm = normalizar(nome)
    if (not principais or len(boas) < 0.7 * len(todos) or com_numero > 1  # lixo / linha de produto
            or RE_ENDERECO.match(norm) or RE_CIDADE_UF.search(norm)):  # endereço / cidade - UF
        return ""
    return _limpar_nome(" ".join(tokens))


def _sufixos(linha: str, linha_norm: str) -> list:
    """LTDA, S.A., ME... (ignora "sa", "me" minúsculos, que costumam ser lixo do OCR)."""
    return [m for m in RE_SUFIXO.finditer(linha_norm) if not linha[m.start():m.end()].islower()]


def _nome_da_linha(linha: str, linha_norm: str) -> str:
    """Corta rótulos do DANFE que o OCR juntou na mesma linha do nome."""
    corte = RE_ROTULOS.search(linha_norm)
    fim = corte.start() if corte else len(linha)
    sufixos = _sufixos(linha[:fim], linha_norm[:fim])
    if sufixos:
        fim = sufixos[-1].end()
    nome = _nome_plausivel(_limpar_nome(linha[:fim]))
    return nome if sum(c.isalpha() for c in nome) >= 4 else ""


def _nome_do_canhoto(texto: str, norm: str) -> str:
    """ "RECEBEMOS DE <FORNECEDOR> OS PRODUTOS..." (aceita erros do OCR nessas palavras)."""
    tokens = list(RE_PALAVRA.finditer(norm))
    for i in range(len(tokens) - 2):
        if tokens[i + 1].group() != "DE" or not _parece_recebemos(tokens[i].group()):
            continue
        for j in range(i + 3, min(i + 30, len(tokens))):
            palavra = re.sub(r"[^A-Z]", "", tokens[j].group())
            if len(palavra) >= 5 and any(_parece(palavra, fim) for fim in FIM_DO_CANHOTO):
                k = j - 1
                if tokens[k].group() in ("OS", "O", "AS", "A"):
                    k -= 1
                if k >= i + 2:
                    nome = _nome_plausivel(_limpar_nome(texto[tokens[i + 2].start():tokens[k].end()]))
                    if sum(c.isalpha() for c in nome) >= 4:
                        return nome
                break
    return ""


def encontrar_fornecedor(texto: str, norm: str) -> str:
    # 1) Canhoto: "RECEBEMOS DE <FORNECEDOR> OS PRODUTOS..."
    nome = _nome_do_canhoto(texto, norm)
    if nome:
        return nome

    linhas, linhas_norm = texto.split("\n"), norm.split("\n")
    inicio_linha = [0]
    for linha in linhas:
        inicio_linha.append(inicio_linha[-1] + len(linha) + 1)

    def linha_de(pos: int) -> int:
        return max(i for i, ini in enumerate(inicio_linha) if ini <= pos)

    # 2) Bloco "IDENTIFICAÇÃO DO EMITENTE" -> próxima linha com um nome
    m = RE_EMITENTE.search(norm)
    if m:
        i = linha_de(m.end())
        resto = texto[m.end():inicio_linha[i + 1] - 1]
        nome = _nome_da_linha(resto, norm[m.end():inicio_linha[i + 1] - 1])
        if nome:
            return nome
        for j in range(i + 1, min(i + 7, len(linhas))):
            if RE_FIM_EMITENTE.search(linhas_norm[j]):
                break
            nome = _nome_da_linha(linhas[j], linhas_norm[j])
            if nome:
                return nome

    # 3) NFS-e: "PRESTADOR DE SERVIÇOS ... Razão Social: <nome>"
    m = RE_PRESTADOR.search(norm)
    if m:
        r = RE_RAZAO_SOCIAL.search(norm, m.end(), m.end() + 400)
        if r:
            j = linha_de(r.end())
            ini = r.end() - inicio_linha[j]
            nome = _nome_da_linha(linhas[j][ini:], linhas_norm[j][ini:])
            if nome:
                return nome

    # 4) Primeira linha com LTDA / S.A. / EIRELI... antes do destinatário
    fim = RE_FIM_EMITENTE.search(norm)
    ultima = linha_de(fim.start()) if fim else len(linhas)
    for j in range(ultima):
        if "RECEBEMOS" in linhas_norm[j] or not _sufixos(linhas[j], linhas_norm[j]):
            continue
        nome = _nome_da_linha(linhas[j], linhas_norm[j])
        if nome:
            return nome
    return ""


def extrair_dados(texto: str, chave_barras: str = "", cadastro: dict | None = None,
                  valor_geometria: str = "") -> DadosNF:
    """Extrai fornecedor, número da NF, pedidos e valor total do texto de uma página."""
    norm = normalizar(texto)
    num = corrigir_digitos(norm)
    dados = DadosNF()
    parcial = ""

    if chave_plausivel(chave_barras):
        dados.chave = chave_barras
    else:
        dados.chave = encontrar_chave(num) or encontrar_chave(norm)

    if dados.chave:
        dados.cnpj = dados.chave[6:20]
        dados.serie = str(int(dados.chave[22:25]))
        dados.numero_nf = str(int(dados.chave[25:34]))
        dados.origem_nf = "chave"
    else:
        dados.numero_nf, dados.origem_nf = encontrar_numero_nf(num)
        dados.serie = encontrar_serie(num)
        dados.cnpj = encontrar_cnpj_emitente(norm)
        # Chave com algum dígito ilegível: o CNPJ dela é confiável; o número precisa de conferência
        parcial = encontrar_chave(num, chave_estrutura_ok)
        if parcial:
            dados.cnpj = parcial[6:20]
            dados.serie = dados.serie or str(int(parcial[22:25]))
            numero_chave = str(int(parcial[25:34]))
            if not dados.numero_nf:
                dados.numero_nf, dados.origem_nf = numero_chave, "chave_parcial"
            elif dados.numero_nf != numero_chave:
                dados.origem_nf = "divergente"

    dados.pedidos = encontrar_pedidos(num, ignorar=dados.chave or parcial)
    if valor_geometria:
        dados.valor_total, dados.origem_valor = valor_geometria, "rotulo"
    else:
        dados.valor_total, dados.origem_valor = encontrar_valor_total(num)

    if cadastro and dados.cnpj in cadastro:
        dados.fornecedor = cadastro[dados.cnpj]
    else:
        dados.fornecedor = encontrar_fornecedor(texto, norm)
    return dados


# ----------------------------------------------------------------------------
# Imagens e OCR
# ----------------------------------------------------------------------------

def configurar_tesseract() -> str:
    """Localiza o Tesseract (inclusive no Windows) e devolve o idioma a usar."""
    import pytesseract

    candidatos = [os.environ.get("TESSERACT_CMD"), shutil.which("tesseract"),
                  r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                  r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                  os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe")]
    for caminho in candidatos:
        if caminho and Path(caminho).is_file():
            pytesseract.pytesseract.tesseract_cmd = caminho
            break
    else:
        sys.exit("ERRO: Tesseract OCR não encontrado.\n"
                 "Instale pelo link https://github.com/UB-Mannheim/tesseract/wiki "
                 "(marque o idioma Portuguese) ou informe o caminho na variável TESSERACT_CMD.")

    if "por" in pytesseract.get_languages(config=""):
        return "por"
    print("AVISO: idioma português ('por') não instalado no Tesseract; usando inglês. "
          "A leitura fica melhor com o português instalado.")
    return "eng"


def carregar_paginas(caminho: Path):
    """Gera (nº da página, imagem, texto embutido) para PDFs e imagens."""
    from PIL import Image, ImageOps, ImageSequence

    if caminho.suffix.lower() == ".pdf":
        import pymupdf as fitz
        with fitz.open(caminho) as doc:
            for n, pagina in enumerate(doc, 1):
                pix = pagina.get_pixmap(dpi=DPI_PDF, colorspace=fitz.csGRAY)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                yield n, img, pagina.get_text()
    else:
        with Image.open(caminho) as im:
            for n, quadro in enumerate(ImageSequence.Iterator(im), 1):
                yield n, ImageOps.exif_transpose(quadro.copy()), ""


def preparar_imagem(img):
    from PIL import Image, ImageOps

    img = img.convert("L")
    maior = max(img.size)
    if maior < 2400:  # digitalização em baixa resolução: amplia para o OCR
        fator = 3000 / maior
        img = img.resize((int(img.width * fator), int(img.height * fator)), Image.LANCZOS)
    elif maior > 4500:
        img.thumbnail((4500, 4500), Image.LANCZOS)
    return ImageOps.autocontrast(img, cutoff=1)


def binarizar(img):
    """Preto e branco com limiar automático (Otsu)."""
    hist = img.histogram()[:256]
    total = sum(hist)
    soma_total = sum(i * h for i, h in enumerate(hist))
    soma_fundo = peso_fundo = 0
    melhor_var, limiar = 0.0, 128
    for t in range(256):
        peso_fundo += hist[t]
        peso_frente = total - peso_fundo
        if peso_fundo == 0:
            continue
        if peso_frente == 0:
            break
        soma_fundo += t * hist[t]
        media_fundo = soma_fundo / peso_fundo
        media_frente = (soma_total - soma_fundo) / peso_frente
        var = peso_fundo * peso_frente * (media_fundo - media_frente) ** 2
        if var > melhor_var:
            melhor_var, limiar = var, t
    return img.point(lambda p: 255 if p > limiar else 0)


def corrigir_rotacao(img):
    """Desvira a folha se ela foi escaneada de lado ou de cabeça para baixo."""
    import pytesseract

    try:
        osd = pytesseract.image_to_osd(img, output_type=pytesseract.Output.DICT)
    except Exception:
        return None
    angulo = osd.get("rotate", 0)
    if angulo and osd.get("orientation_conf", 0) >= 1.5:
        return img.rotate(-angulo, expand=True)
    return None


def ler_codigo_barras(img) -> str:
    """Lê a chave de acesso pelo código de barras do DANFE (se o pyzbar estiver instalado)."""
    try:
        from pyzbar import pyzbar
        codigos = pyzbar.decode(img)
    except Exception:
        return ""
    for codigo in codigos:
        chave = encontrar_chave(codigo.data.decode("ascii", "ignore").upper())
        if chave:
            return chave
    return ""


def ocr(img, idioma: str, psm: int) -> tuple[str, list]:
    """Texto (linha a linha) e palavras com a posição: [texto, x, y, largura, altura, linha]."""
    import pytesseract
    d = pytesseract.image_to_data(img, lang=idioma, config=f"--psm {psm}",
                                  output_type=pytesseract.Output.DICT)
    linhas: dict = {}
    palavras = []
    for i, texto in enumerate(d["text"]):
        if not texto.strip():
            continue
        linha = d["block_num"][i] * 1000000 + d["par_num"][i] * 1000 + d["line_num"][i]
        linhas.setdefault(linha, []).append(texto)
        palavras.append([texto, d["left"][i], d["top"][i], d["width"][i], d["height"][i], linha])
    return "\n".join(" ".join(p) for p in linhas.values()), palavras


def recortar_topo(img, fracao: float = 0.6, zoom: float = 2.0):
    """Parte de cima da folha (número, chave, fornecedor, valores), ampliada."""
    from PIL import Image
    topo = img.crop((0, 0, img.width, int(img.height * fracao)))
    return topo.resize((int(topo.width * zoom), int(topo.height * zoom)), Image.LANCZOS)


def ler_pagina(img, texto_embutido: str, idioma: str, cadastro: dict) -> tuple[DadosNF, str, list]:
    img = preparar_imagem(img)
    chave_barras = ler_codigo_barras(img)
    base = [texto_embutido] if len(texto_embutido.strip()) >= 100 else []  # PDF com texto
    leituras: list = []  # (texto, palavras)

    def montar() -> tuple[DadosNF, str]:
        textos = base + [t for t, _ in leituras] + [texto_visual(p) for _, p in leituras if p]
        valor = next((v for v in (valor_pelo_rotulo(p) for _, p in leituras) if v), "")
        return extrair_dados("\n".join(textos), chave_barras, cadastro, valor), "\n\n".join(textos)

    dados, texto = montar()
    if not dados.completo:
        leituras.append(ocr(img, idioma, psm=3))
        dados, texto = montar()
    if not dados.completo:
        girada = corrigir_rotacao(img)
        if girada is not None:  # folha virada: o texto lido antes era lixo
            img = girada
            leituras = [ocr(img, idioma, psm=3)]
            dados, texto = montar()
    if not dados.completo:
        leituras.append(ocr(recortar_topo(img), idioma, psm=11))
        dados, texto = montar()
    if not dados.completo:
        leituras.append(ocr(binarizar(img), idioma, psm=6))
        dados, texto = montar()
    return dados, texto, leituras


def processar_arquivo(caminho: Path, nome: str, idioma: str, cadastro: dict,
                      pasta_textos: Path | None) -> list[RegistroNF]:
    registros: list[RegistroNF] = []
    for n, img, texto_embutido in carregar_paginas(caminho):
        dados, texto, leituras = ler_pagina(img, texto_embutido, idioma, cadastro)
        if pasta_textos:
            base = re.sub(r"[\\/:]", "_", nome)
            (pasta_textos / f"{base}_pag{n}.txt").write_text(texto, encoding="utf-8")
            (pasta_textos / f"{base}_pag{n}.json").write_text(
                json.dumps([{"palavras": p} for _, p in leituras], ensure_ascii=False), encoding="utf-8")

        if dados.vazio:
            if len(re.sub(r"\s", "", texto)) >= 30:  # ignora folhas em branco (verso)
                registros.append(RegistroNF(nome, [n], dados))
            continue

        # A mesma NF em várias páginas vira uma linha só
        for reg in registros:
            mesmo_emitente = not (reg.dados.cnpj and dados.cnpj) or reg.dados.cnpj == dados.cnpj
            if dados.numero_nf and reg.dados.numero_nf == dados.numero_nf and mesmo_emitente:
                reg.paginas.append(n)
                reg.dados.completar_com(dados)
                break
        else:
            registros.append(RegistroNF(nome, [n], dados))
    return registros


# ----------------------------------------------------------------------------
# Cadastro de fornecedores e planilha
# ----------------------------------------------------------------------------

def carregar_cadastro(caminho: Path) -> dict[str, str]:
    """Lê o CSV opcional 'CNPJ;Fornecedor' para padronizar os nomes."""
    if not caminho.is_file():
        return {}
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            conteudo = caminho.read_text(encoding=codificacao)
            break
        except UnicodeDecodeError:
            continue
    linhas = conteudo.splitlines()
    separador = ";" if linhas and ";" in linhas[0] else ","
    cadastro = {}
    for partes in csv.reader(linhas, delimiter=separador):
        if len(partes) < 2:
            continue
        cnpj = re.sub(r"[^0-9A-Z]", "", partes[0].upper())
        if len(cnpj) == 14 and partes[1].strip():
            cadastro[cnpj] = partes[1].strip()
    return cadastro


def _tem_sufixo(nome: str) -> bool:
    return bool(RE_SUFIXO_FINAL.search(normalizar(nome)))


def _semelhanca(a: str, b: str, pelo_fim: bool = False) -> float:
    """0 a 1. Um nome cortado pelo OCR também é comparado com o começo (ou o fim) do outro."""
    a, b = normalizar(a), normalizar(b)
    if not a or not b:
        return 0.0
    s = 1 - distancia(a, b) / max(len(a), len(b))
    if 12 <= len(a) < len(b):
        s = max(s, 1 - distancia(a, b[:len(a)]) / len(a))
        if pelo_fim:
            s = max(s, 1 - distancia(a, b[-len(a):]) / len(a))
    return s


def padronizar_fornecedores(registros: list[RegistroNF]) -> None:
    """O mesmo fornecedor lido de jeitos diferentes em NFs diferentes passa a ter um nome só."""
    def contar():
        contagem: Counter = Counter()
        primeira: dict[str, int] = {}
        for i, r in enumerate(registros):
            if r.dados.fornecedor:
                contagem[r.dados.fornecedor] += 1
                primeira.setdefault(r.dados.fornecedor, i)
        return lambda n: (_tem_sufixo(n), contagem[n], len(n), -primeira[n]), contagem

    # 1) Nome cortado ou mal lido (sem LTDA/S.A.) parecido com um nome completo de outra NF
    ranking, contagem = contar()
    completos = sorted((n for n in contagem if _tem_sufixo(n)), key=ranking, reverse=True)
    for r in registros:
        nome = r.dados.fornecedor
        if nome and not _tem_sufixo(nome):
            for completo in completos:
                if _semelhanca(nome, completo) >= 0.8:
                    r.dados.fornecedor = completo
                    break

    # 2) Mesmo CNPJ = mesmo fornecedor: usa a melhor leitura do grupo
    ranking, _ = contar()
    grupos: dict[str, list[RegistroNF]] = {}
    for r in registros:
        if r.dados.cnpj:
            grupos.setdefault(r.dados.cnpj, []).append(r)
    for grupo in grupos.values():
        nomes = [r.dados.fornecedor for r in grupo if r.dados.fornecedor]
        if not nomes:
            continue
        melhor = max(nomes, key=ranking)
        for r in grupo:
            nome = r.dados.fornecedor
            if not nome or not _tem_sufixo(nome) or _semelhanca(nome, melhor, pelo_fim=True) >= 0.8:
                r.dados.fornecedor = melhor


def marcar_avisos(registros: list[RegistroNF]) -> None:
    chaves = Counter((r.dados.cnpj or r.dados.fornecedor, r.dados.numero_nf)
                     for r in registros if r.dados.numero_nf)
    for r in registros:
        d = r.dados
        if r.avisos:  # erro ao abrir o arquivo
            continue
        if d.vazio:
            r.avisos.append("nada identificado na página")
            continue
        if not d.fornecedor:
            r.avisos.append("fornecedor não encontrado")
        if not d.numero_nf:
            r.avisos.append("nº da NF não encontrado")
        elif d.origem_nf == "texto":
            r.avisos.append("nº da NF lido fora do padrão DANFE")
        elif d.origem_nf == "chave_parcial":
            r.avisos.append("nº da NF lido da chave de acesso com dígito ilegível")
        elif d.origem_nf == "divergente":
            r.avisos.append("nº impresso diferente do nº da chave de acesso")
        elif chaves[(d.cnpj or d.fornecedor, d.numero_nf)] > 1:
            r.avisos.append("NF repetida em outro arquivo")
        if not d.valor_total:
            r.avisos.append("valor total não encontrado")
        elif d.origem_valor == "maior":
            r.avisos.append("valor total estimado (maior valor da nota)")
        if not d.pedidos:
            r.avisos.append("pedido não encontrado")
        elif len(d.pedidos) > 1:
            r.avisos.append("mais de um pedido")


def gerar_planilha(registros: list[RegistroNF], saida: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    colunas = [("Fornecedor", 45), ("Nº NF", 12), ("Pedido", 14), ("Valor Total", 14), ("Série", 7),
               ("CNPJ Fornecedor", 20), ("Chave de Acesso", 48), ("Arquivo", 30),
               ("Página", 8), ("Status", 55)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Notas Fiscais"
    ws.append([nome for nome, _ in colunas])
    for i, (_, largura) in enumerate(colunas, 1):
        cel = ws.cell(row=1, column=i)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="1F4E78")
        cel.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cel.column_letter].width = largura

    amarelo = PatternFill("solid", fgColor="FFF2CC")
    for r in registros:
        d = r.dados
        numero = int(d.numero_nf) if d.numero_nf else ""
        if len(d.pedidos) == 1:
            pedido = int(d.pedidos[0])
        else:
            pedido = " / ".join(d.pedidos)
        status = "OK" if not r.avisos else "CONFERIR: " + "; ".join(r.avisos)
        valor = float(d.valor_total) if d.valor_total else ""
        ws.append([d.fornecedor, numero, pedido, valor, d.serie, formatar_cnpj(d.cnpj), d.chave,
                   r.arquivo, ", ".join(map(str, r.paginas)), status])
        linha = ws.max_row
        for col in (2, 3):
            ws.cell(row=linha, column=col).number_format = "0"
        ws.cell(row=linha, column=4).number_format = "#,##0.00"
        if r.avisos:
            for cel in ws[linha]:
                cel.fill = amarelo

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    try:
        wb.save(saida)
    except PermissionError:  # planilha aberta no Excel
        saida = saida.with_name(f"{saida.stem}_{datetime.now():%Y%m%d_%H%M%S}{saida.suffix}")
        wb.save(saida)
    return saida


# ----------------------------------------------------------------------------
# Programa principal
# ----------------------------------------------------------------------------

def verificar_dependencias() -> None:
    faltando = []
    for modulo, pacote in (("PIL", "pillow"), ("openpyxl", "openpyxl"),
                           ("pytesseract", "pytesseract"), ("pymupdf", "pymupdf")):
        try:
            __import__(modulo)
        except ImportError:
            faltando.append(pacote)
    if faltando:
        sys.exit("ERRO: faltam bibliotecas Python. Rode:\n"
                 f"    pip install {' '.join(faltando)}\n"
                 "ou: pip install -r requirements.txt")


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")

    parser = argparse.ArgumentParser(
        description="Lê NFs escaneadas e gera uma planilha com Fornecedor, Nº NF e Pedido.")
    parser.add_argument("pasta", nargs="?", default=str(PASTA_PADRAO),
                        help="pasta com os PDFs/imagens escaneados (padrão: ./notas)")
    parser.add_argument("-o", "--saida", default=str(SAIDA_PADRAO),
                        help="planilha a gerar (padrão: planilha_nfs.xlsx)")
    parser.add_argument("--cadastro", default=str(CADASTRO_PADRAO),
                        help="CSV opcional com CNPJ;Fornecedor (padrão: fornecedores.csv)")
    parser.add_argument("--salvar-texto", action="store_true",
                        help="salva o texto lido de cada página na pasta textos_ocr")
    args = parser.parse_args(argv)

    pasta = Path(args.pasta)
    if not pasta.is_dir():
        if pasta == PASTA_PADRAO:
            pasta.mkdir(parents=True)
            print(f"Criei a pasta {pasta}. Coloque as NFs escaneadas nela e rode de novo.")
            return 1
        print(f"ERRO: pasta não encontrada: {pasta}")
        return 1

    arquivos = sorted(p for p in pasta.rglob("*") if p.is_file() and p.suffix.lower() in EXTENSOES)
    if not arquivos:
        print(f"Nenhum PDF ou imagem encontrado em {pasta}")
        return 1

    verificar_dependencias()
    idioma = configurar_tesseract()
    cadastro = carregar_cadastro(Path(args.cadastro))
    if cadastro:
        print(f"Cadastro de fornecedores: {len(cadastro)} CNPJs")

    saida = Path(args.saida)
    saida.parent.mkdir(parents=True, exist_ok=True)
    pasta_textos = None
    if args.salvar_texto:
        pasta_textos = saida.parent / "textos_ocr"
        pasta_textos.mkdir(parents=True, exist_ok=True)

    registros: list[RegistroNF] = []
    for i, arquivo in enumerate(arquivos, 1):
        nome = str(arquivo.relative_to(pasta))
        print(f"[{i}/{len(arquivos)}] {nome}")
        try:
            novos = processar_arquivo(arquivo, nome, idioma, cadastro, pasta_textos)
        except Exception as erro:
            print(f"    ERRO ao ler o arquivo: {erro}")
            registros.append(RegistroNF(nome, [], DadosNF(), [f"erro ao ler o arquivo: {erro}"]))
            continue
        for r in novos:
            d = r.dados
            print(f"    NF {d.numero_nf or '?':>9} | {d.fornecedor or '?'} | "
                  f"Pedido {' / '.join(d.pedidos) or '?'} | Valor {d.valor_total or '?'}")
        registros.extend(novos)

    padronizar_fornecedores(registros)
    marcar_avisos(registros)
    caminho = gerar_planilha(registros, saida)
    conferir = sum(1 for r in registros if r.avisos)
    print(f"\nPlanilha gerada: {caminho}")
    print(f"{len(registros)} nota(s) | {len(registros) - conferir} OK | {conferir} para conferir")
    return 0


if __name__ == "__main__":
    sys.exit(main())
