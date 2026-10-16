"""Testes da leitura do texto das NFs (não precisam do Tesseract).

Rodar:  python -m unittest test_robo_nf.py
"""

import random
import tempfile
import unittest
from pathlib import Path

import robo_nf as r


def montar_chave(cnpj, numero, serie=1, uf="35", aamm="2610"):
    base = f"{uf}{aamm}{cnpj}55{serie:03d}{numero:09d}1{random.randint(0, 99999999):08d}"
    return base + str(r.dv_mod11(base[::-1], [2, 3, 4, 5, 6, 7, 8, 9] * 6))


def agrupar(chave):
    return " ".join(chave[i:i + 4] for i in range(0, 44, 4))


CNPJ = "12345678000195"

DANFE = """RECEBEMOS DE FORNECEDOR EXEMPLO INDÚSTRIA E COMÉRCIO LTDA OS PRODUTOS E/OU SERVIÇOS
CONSTANTES DA NOTA FISCAL ELETRÔNICA INDICADA AO LADO          NF-e
DATA DE RECEBIMENTO  IDENTIFICAÇÃO E ASSINATURA DO RECEBEDOR   Nº 000.012.345
                                                               SÉRIE 001
IDENTIFICAÇÃO DO EMITENTE
FORNECEDOR EXEMPLO INDÚSTRIA E COMÉRCIO LTDA       DANFE
RUA DAS INDÚSTRIAS, 1500                           Nº 000.012.345  SÉRIE 001
CHAVE DE ACESSO
{chave}
INSCRIÇÃO ESTADUAL 123.456.789.110    CNPJ 12.345.678/0001-95
DESTINATÁRIO / REMETENTE
NOME / RAZÃO SOCIAL
MINHA EMPRESA COMPRADORA S.A.        CNPJ 11.222.333/0001-81
DADOS ADICIONAIS
Pedido de compra: 4500012345   Valor aprox. tributos R$ 45,32
"""


class TestChave(unittest.TestCase):
    def test_chave_agrupada_e_validada(self):
        chave = montar_chave(CNPJ, 12345)
        self.assertTrue(r.chave_plausivel(chave))
        self.assertEqual(r.encontrar_chave(f"CHAVE DE ACESSO\n{agrupar(chave)}\n"), chave)

    def test_chave_com_digito_errado_e_ignorada(self):
        chave = montar_chave(CNPJ, 12345)
        errada = chave[:-1] + str((int(chave[-1]) + 1) % 10)
        self.assertFalse(r.chave_plausivel(errada))
        self.assertEqual(r.encontrar_chave(agrupar(errada)), "")

    def test_cnpj_alfanumerico(self):
        # Exemplo oficial da Receita Federal para o CNPJ alfanumérico
        self.assertTrue(r.cnpj_valido("12ABC34501DE35"))
        chave = montar_chave("12ABC34501DE35", 777)
        self.assertTrue(r.chave_plausivel(chave))
        dados = r.extrair_dados(f"CHAVE DE ACESSO {agrupar(chave)}\nPEDIDO 4500000001")
        self.assertEqual(dados.numero_nf, "777")
        self.assertEqual(dados.cnpj, "12ABC34501DE35")


class TestExtracao(unittest.TestCase):
    def test_danfe_completo(self):
        chave = montar_chave(CNPJ, 12345)
        dados = r.extrair_dados(DANFE.format(chave=agrupar(chave)))
        self.assertEqual(dados.fornecedor, "FORNECEDOR EXEMPLO INDÚSTRIA E COMÉRCIO LTDA")
        self.assertEqual(dados.numero_nf, "12345")
        self.assertEqual(dados.serie, "1")
        self.assertEqual(dados.cnpj, CNPJ)
        self.assertEqual(dados.pedidos, ["4500012345"])
        self.assertEqual(dados.origem_nf, "chave")

    def test_danfe_sem_chave_usa_numero_impresso(self):
        dados = r.extrair_dados(DANFE.format(chave="(ilegível)"))
        self.assertEqual(dados.numero_nf, "12345")
        self.assertEqual(dados.origem_nf, "danfe")
        self.assertEqual(dados.serie, "1")
        self.assertEqual(dados.cnpj, CNPJ)  # do emitente, não do destinatário

    def test_chave_do_codigo_de_barras_tem_prioridade(self):
        chave = montar_chave(CNPJ, 999)
        dados = r.extrair_dados(DANFE.format(chave=""), chave_barras=chave)
        self.assertEqual(dados.numero_nf, "999")

    def test_numero_com_letras_trocadas_pelo_ocr(self):
        dados = r.extrair_dados("DANFE\nNº 000.0I2.345 SÉRIE 1\nPEDIDO 4500012345")
        self.assertEqual(dados.numero_nf, "12345")


class TestPedido(unittest.TestCase):
    def test_varios_pedidos_na_ordem(self):
        texto = "PEDIDOS: 4500099887, 4500099888 e novamente 4500099887"
        self.assertEqual(r.encontrar_pedidos(texto), ["4500099887", "4500099888"])

    def test_pedido_quebrado_pelo_ocr(self):
        num = r.corrigir_digitos(r.normalizar("Seu pedido nº 4500 0I2345"))
        self.assertEqual(r.encontrar_pedidos(num), ["4500012345"])
        num = r.corrigir_digitos(r.normalizar("Ped. compra: 45OOO12345"))
        self.assertEqual(r.encontrar_pedidos(num), ["4500012345"])

    def test_nao_confunde_com_outros_numeros(self):
        texto = ("CNPJ 45.123.456/0001-00  CNPJ 45123456000100  FONE (45) 3333-4444\n"
                 "EAN 7894500012345  VALOR 4.500.123,45  PROTOCOLO 145000123456789\n")
        self.assertEqual(r.encontrar_pedidos(texto), [])

    def test_ignora_digitos_da_chave(self):
        chave = "35261045000123450001550010000012341000000017"
        self.assertEqual(r.encontrar_pedidos(f"CHAVE {chave}", ignorar=chave), [])
        self.assertEqual(r.encontrar_pedidos(agrupar(chave), ignorar=chave), [])

    def test_pedido_com_11_digitos_nao_conta(self):
        self.assertEqual(r.encontrar_pedidos("PEDIDO 45000123456"), [])


class TestFornecedor(unittest.TestCase):
    def fornecedor(self, texto):
        return r.encontrar_fornecedor(texto, r.normalizar(texto))

    def test_bloco_do_emitente_sem_canhoto(self):
        texto = ("IDENTIFICAÇÃO DO EMITENTE\n\nMETALÚRGICA SILVA S/A          DANFE\n"
                 "AV. BRASIL, 100\nDESTINATÁRIO\nMINHA EMPRESA LTDA")
        self.assertEqual(self.fornecedor(texto), "METALÚRGICA SILVA S/A")

    def test_linha_com_ltda_antes_do_destinatario(self):
        texto = ("DANFE  0-ENTRADA 1-SAÍDA\nPAPELARIA CENTRAL LTDA - ME   CNPJ 11.111.111/0001-11\n"
                 "DESTINATÁRIO / REMETENTE\nMINHA EMPRESA COMPRADORA LTDA")
        self.assertEqual(self.fornecedor(texto), "PAPELARIA CENTRAL LTDA - ME")

    def test_nfse_prestador(self):
        texto = ("PREFEITURA MUNICIPAL\nNÚMERO DA NOTA: 1234\nPRESTADOR DE SERVIÇOS\n"
                 "Razão Social: SERVIÇOS TÉCNICOS ALFA LTDA\nTOMADOR DE SERVIÇOS\n"
                 "Razão Social: MINHA EMPRESA S.A.\nPedido 4500001111")
        dados = r.extrair_dados(texto)
        self.assertEqual(dados.fornecedor, "SERVIÇOS TÉCNICOS ALFA LTDA")
        self.assertEqual(dados.numero_nf, "1234")
        self.assertEqual(dados.origem_nf, "texto")
        self.assertEqual(dados.pedidos, ["4500001111"])

    def test_cadastro_padroniza_nome(self):
        chave = montar_chave(CNPJ, 12345)
        dados = r.extrair_dados(DANFE.format(chave=agrupar(chave)),
                                cadastro={CNPJ: "Fornecedor Exemplo (cód. SAP 100200)"})
        self.assertEqual(dados.fornecedor, "Fornecedor Exemplo (cód. SAP 100200)")


class TestCasosReais(unittest.TestCase):
    """Erros vistos nas primeiras NFs reais (lidas com o OCR do Windows)."""

    def test_inscricao_estadual_nao_vira_numero_da_nf(self):
        texto = "INSCRIÇÃO ESTADUAL\n252.405.749\nCNPJ 05.283.785/0001-70\nPEDIDO 4501216027"
        self.assertEqual(r.extrair_dados(texto).numero_nf, "")
        texto = "INSCRIÇÃO ESTADUAL 252.405.749   NF-e Nº 000.093.269 SÉRIE 1"
        self.assertEqual(r.extrair_dados(texto).numero_nf, "93269")

    def test_numero_sem_rotulo_precisa_de_zeros_a_esquerda(self):
        self.assertEqual(r.encontrar_numero_nf("DANFE 252.405.749"), ("", ""))
        self.assertEqual(r.encontrar_numero_nf("DANFE 000.093.269"), ("93269", "danfe"))

    def test_espacos_do_ocr_no_numero(self):
        num = r.corrigir_digitos(r.normalizar("Nº 000. 093 .269"))
        self.assertEqual(r.encontrar_numero_nf(num), ("93269", "danfe"))

    def test_numero_sem_pontos_antes_da_serie(self):
        for texto in ("NF-e\nNº 93269\nSÉRIE 1", "Nº. 000093269 SÉRIE: 001"):
            num = r.corrigir_digitos(r.normalizar(texto))
            self.assertEqual(r.encontrar_numero_nf(num), ("93269", "danfe"), texto)

    def test_chave_quebrada_em_duas_linhas(self):
        grupos = agrupar(montar_chave(CNPJ, 85482))
        texto = f"CHAVE DE ACESSO\n{grupos[:27]}\n{grupos[27:]}\nConsulta de autenticidade"
        self.assertEqual(r.extrair_dados(texto).numero_nf, "85482")

    def test_chave_com_digito_ilegivel(self):
        chave = montar_chave(CNPJ, 85482)
        errada = next(e for e in (chave[:40] + str(d) + chave[41:] for d in range(10))
                      if not r.chave_plausivel(e))
        dados = r.extrair_dados(f"CHAVE DE ACESSO {agrupar(errada)}")
        self.assertEqual((dados.numero_nf, dados.origem_nf, dados.cnpj), ("85482", "chave_parcial", CNPJ))
        self.assertEqual(dados.chave, "")  # chave errada não vai para a planilha
        dados = r.extrair_dados(f"Nº 000.085.482 SÉRIE 1\nCHAVE DE ACESSO {agrupar(errada)}")
        self.assertEqual((dados.numero_nf, dados.origem_nf), ("85482", "danfe"))
        dados = r.extrair_dados(f"Nº 000.085.483 SÉRIE 1\nCHAVE DE ACESSO {agrupar(errada)}")
        self.assertEqual((dados.numero_nf, dados.origem_nf), ("85483", "divergente"))

    def test_lixo_do_ocr_nao_vira_fornecedor(self):
        for lixo in ("trtVtE a 1.1-snrv oagoav sa", "eoyssuw oavt ov sa", "•v•s sa", "IV ayJ.sa",
                     "svax•UN00Nd a sa.", "bd.Remget•gçi0 Ltda", "van wu.st501 SA.",
                     "OL18LSta•a-m ¿•z sa"):
            for texto in (f"{lixo}\nDESTINATÁRIO / REMETENTE\nMINHA EMPRESA LTDA",
                          f"IDENTIFICAÇÃO DO EMITENTE\n{lixo}\nDESTINATÁRIO / REMETENTE"):
                self.assertEqual(r.encontrar_fornecedor(texto, r.normalizar(texto)), "", texto)

    def test_nomes_reais_continuam_aceitos(self):
        for nome in ("DAL PINO INDUSTRIA DE SERRAS LTDA", "Frigostrella do Brasil Ind.Refrigeracao Ltda",
                     "BETTCHER DO BRASIL COMERCIO DE MAQUINAS LTDA", "METALÚRGICA SILVA S/A",
                     "PAPELARIA CENTRAL LTDA - ME", "JBT MAREL BRASIL LTDA", "FORTUNCERES S.A."):
            self.assertEqual(r._nome_plausivel(nome), nome)

    def test_endereco_cidade_e_produto_nao_viram_fornecedor(self):
        for linha in ("RUA DAS INDÚSTRIAS, 1500 - DISTRITO INDUSTRIAL", "AV. BRASIL, 100",
                      "SÃO PAULO - SP CEP 01000-000", "000123 PARAFUSO SEXTAVADO M8 73181500 100 1,50",
                      "ossaovadaavHS — WIVHTOS TOO IIHAS"):
            texto = f"IDENTIFICAÇÃO DO EMITENTE\n{linha}\nDESTINATÁRIO / REMETENTE"
            self.assertEqual(r.encontrar_fornecedor(texto, r.normalizar(texto)), "", linha)

    def test_tira_pedaco_ilegivel_do_nome(self):
        texto = "RECEBEMOS DE sn«ox COMERCIO DE VALVULAS LTDA OS PRODUTOS CONSTANTES DA NOTA"
        self.assertEqual(r.encontrar_fornecedor(texto, r.normalizar(texto)), "COMERCIO DE VALVULAS LTDA")

    def test_fornecedor_pelo_cnpj_de_outra_nf(self):
        com_nome = r.RegistroNF("20.pdf", [1], r.DadosNF("DIMENSIONAL BRASIL LTDA", "06913480002616", "1"))
        sem_nome = r.RegistroNF("19.pdf", [1], r.DadosNF("", "06913480002616", "2"))
        r.padronizar_fornecedores([sem_nome, com_nome])
        self.assertEqual(sem_nome.dados.fornecedor, "DIMENSIONAL BRASIL LTDA")



def palavra(texto, x, y, largura=None, altura=20):
    return [texto, x, y, largura if largura is not None else 11 * len(texto), altura]


# Quadro "CÁLCULO DO IMPOSTO" de um DANFE, como o OCR devolve (palavras e posições)
QUADRO_IMPOSTO = [
    palavra("VALOR", 600, 40), palavra("TOTAL", 670, 40), palavra("DOS", 735, 40), palavra("PRODUTOS", 775, 40),
    palavra("2.400,00", 800, 70),
    palavra("VALOR", 300, 100), palavra("DO", 365, 100), palavra("IPI", 395, 100),
    palavra("VALOR", 600, 100), palavra("TOTAL", 670, 100), palavra("DA", 735, 100), palavra("NOTA", 765, 100),
    palavra("120,00", 380, 130), palavra("2.520,00", 800, 130),
]

# Caso real do teste: o valor do IPI (coluna vizinha) mais perto do rótulo do que o total
QUADRO_IPI_PERTO = [
    ["VALOR", 1395, 1037, 46, 21], ["DO", 1450, 1037, 30, 21], ["IPI.", 1490, 1037, 40, 21],
    ["VALOR", 1660, 1042, 90, 19], ["TOTALDA", 1758, 1042, 98, 18], ["NOTA", 1868, 1042, 56, 19],
    ["120,00", 1554, 1078, 99, 28], ["4.025,00", 1926, 1079, 130, 33],
]


class TestLeituraMelhor(unittest.TestCase):
    """Zeros lidos como letra O, número sem "Nº", valor total, canhoto com erro de OCR."""

    def test_zeros_lidos_como_letra(self):
        num = r.corrigir_digitos(r.normalizar("Nº OOO.O93.269  SC-401  LTDA  ISO  S/A"))
        self.assertIn("000.093.269", num)
        self.assertIn("SC-401  LTDA  ISO  S/A", num)
        self.assertEqual(r.encontrar_numero_nf(num), ("93269", "danfe"))

    def test_numero_com_virgula_espacos_e_q(self):
        for texto in ("Nº 000,093,269", "N° 000 093 269", "NQ 000.093.269", "000093269\nSÉRIE 1"):
            num = r.corrigir_digitos(r.normalizar(texto))
            self.assertEqual(r.encontrar_numero_nf(num), ("93269", "danfe"), texto)

    def test_valor_total_na_mesma_linha(self):
        for texto in ("VALOR TOTAL DA NOTA R$ 1.234,56", "V. TOTAL DA NOTA: 1234,56", "Valor total da NF-e 1.234,56",
                      "VALOR TOTALDA NOTA 1.234,56", "VALOR TOTAL. DA NOTA 1.234,56"):
            num = r.corrigir_digitos(r.normalizar(texto))
            self.assertEqual(r.encontrar_valor_total(num), ("1234.56", "rotulo"), texto)

    def test_valor_total_na_linha_de_baixo(self):
        texto = ("VALOR DO FRETE VALOR DO SEGURO DESCONTO OUTRAS DESP. VALOR DO IPI VALOR TOTAL DA NOTA\n"
                 "0,00 0,00 0,00 0,00 120,00 2.520,00")
        self.assertEqual(r.encontrar_valor_total(r.normalizar(texto)), ("2520.00", "rotulo"))

    def test_valor_total_estimado_sem_rotulo(self):
        texto = "VALOR TOTAL DOS PRODUTOS 2.400,00\nTRIBUTOS R$ 45,32 (19,70%)\nPESO 3.500,000"
        self.assertEqual(r.encontrar_valor_total(r.normalizar(texto)), ("2400.00", "maior"))

    def test_valor_total_pela_posicao_na_folha(self):
        self.assertEqual(r.valor_pelo_rotulo(QUADRO_IMPOSTO), "2520.00")
        self.assertEqual(r.valor_pelo_rotulo(QUADRO_IPI_PERTO), "4025.00")
        visual = r.texto_visual(QUADRO_IMPOSTO).split("\n")
        self.assertEqual(visual[2], "VALOR DO IPI | VALOR TOTAL DA NOTA")
        self.assertEqual(visual[3], "120,00 | 2.520,00")
        dados = r.extrair_dados(r.texto_visual(QUADRO_IMPOSTO))
        self.assertEqual((dados.valor_total, dados.origem_valor), ("2520.00", "rotulo"))

    def test_folha_torta(self):
        # 0,6 grau de inclinação: o fim da linha fica mais alto que o começo
        frase = "RECEBEMOS DE DISTRIBUIDORA BOA VISTA EIRELI OS PRODUTOS E/OU SERVIÇOS".split()
        palavras, x = [], 60
        for k, t in enumerate(frase):
            palavras.append([t, x, 100 - 0.0105 * x, 11 * len(t), 24, 1])
            x += 11 * len(t) + 12
        palavras.insert(0, ["CONSTANTES", x, 100 - 0.0105 * x - 9, 110, 18, 1])  # caixa mais alta (caso real)
        palavras += [["NOTA", 60, 128, 44, 24, 2], ["FISCAL", 116, 128, 66, 24, 2]]
        visual = r.texto_visual(palavras).split("\n")
        self.assertEqual(visual[0], " ".join(frase) + " CONSTANTES")
        self.assertEqual(r.encontrar_fornecedor(visual[0], r.normalizar(visual[0])), "DISTRIBUIDORA BOA VISTA EIRELI")

    def test_canhoto_com_erro_de_ocr(self):
        for texto, nome in (
                ("RFEBEMOS DE DIMENSIONAL BRASIL SOLUCOES LTDA OS PRODUTOS E/OU SERVIÇOS", "DIMENSIONAL BRASIL SOLUCOES LTDA"),
                ("RECEBEM(S DE DAL PINO INDUSTRIA DE SERRAS LTDA\nOS PRODUT0S CONSTANTES", "DAL PINO INDUSTRIA DE SERRAS LTDA")):
            self.assertEqual(r.encontrar_fornecedor(texto, r.normalizar(texto)), nome)
        self.assertEqual(r._nome_plausivel("DE JARVIS DO BRASIL FERRAMENTAS INDUSTRIAIS LTDA"),
                         "JARVIS DO BRASIL FERRAMENTAS INDUSTRIAIS LTDA")
        self.assertEqual(r._nome_plausivel("COMERCIAL DE FERRAGENS NOVA ERA auxiliar da"),
                         "COMERCIAL DE FERRAGENS NOVA ERA")

    def test_padroniza_grafias_do_mesmo_fornecedor(self):
        certo = "JARVIS DO BRASIL FERRAMENTAS INDUSTRIAIS LTDA"
        nomes = [(certo, ""), (certo, ""), ("JARVIS DO BRASL", "46537536000141"), ("JARVIS DO BRASIL FERRAMTXMS", ""),
                 ("JARVLS DO BRASIL FERRAP,mT4S", ""), ("NDUSIRIAIS LTDA", "46537536000141"),
                 ("BETTCHER DO BRASIL COMERCIO DE MAQUINAS LTDA", ""), ("DAL PENO INDUSTRIA DE SERRAS LTDA", ""),
                 ("COMERCIAL ELETRICA PJ LTDA", ""), ("COMERCIAL ELETRICA SC LTDA", "")]
        registros = [r.RegistroNF(f"{i}.pdf", [1], r.DadosNF(n, c)) for i, (n, c) in enumerate(nomes)]
        r.padronizar_fornecedores(registros)
        self.assertEqual([x.dados.fornecedor for x in registros],
                         [certo] * 6 + [n for n, _ in nomes[6:]])


class TestPlanilha(unittest.TestCase):
    def test_cadastro_csv_excel_brasileiro(self):
        with tempfile.TemporaryDirectory() as pasta:
            csv_path = Path(pasta) / "fornecedores.csv"
            csv_path.write_bytes("CNPJ;Fornecedor\n12.345.678/0001-95;Fornecedor Exemplo Ind. e Com.\n"
                                 .encode("cp1252"))
            self.assertEqual(r.carregar_cadastro(csv_path), {CNPJ: "Fornecedor Exemplo Ind. e Com."})

    def test_gera_planilha_com_avisos(self):
        from openpyxl import load_workbook

        ok = r.RegistroNF("a.pdf", [1], r.DadosNF("FORN A LTDA", CNPJ, "10", "1", ["4500000001"], "", "danfe",
                                                  "2520.00", "rotulo"))
        dois = r.RegistroNF("b.pdf", [1, 2], r.DadosNF("FORN B LTDA", "", "20", "", ["4500000002", "4500000003"]))
        sem_pedido = r.RegistroNF("c.jpg", [1], r.DadosNF("FORN C LTDA", "", "30"))
        repetida = r.RegistroNF("d.jpg", [1], r.DadosNF("FORN A LTDA", CNPJ, "10", "1", ["4500000001"], "", "danfe"))
        registros = [ok, dois, sem_pedido, repetida]
        r.marcar_avisos(registros)
        self.assertIn("mais de um pedido", dois.avisos)
        self.assertIn("pedido não encontrado", sem_pedido.avisos)
        self.assertIn("NF repetida em outro arquivo", repetida.avisos)
        self.assertIn("valor total não encontrado", repetida.avisos)

        with tempfile.TemporaryDirectory() as pasta:
            saida = r.gerar_planilha(registros, Path(pasta) / "nfs.xlsx")
            linhas = list(load_workbook(saida).active.iter_rows(values_only=True))
        self.assertEqual(linhas[0][:4], ("Fornecedor", "Nº NF", "Pedido", "Valor Total"))
        self.assertEqual(linhas[1][:4], ("FORN A LTDA", 10, 4500000001, 2520.0))
        self.assertEqual(linhas[2][2], "4500000002 / 4500000003")
        self.assertEqual(linhas[2][8], "1, 2")
        self.assertTrue(linhas[3][9].startswith("CONFERIR"))


if __name__ == "__main__":
    unittest.main()
